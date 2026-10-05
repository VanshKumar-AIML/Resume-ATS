"""
file_extractor.py
------------------
Extracts plain text from an uploaded resume file (PDF or DOCX).

Pipeline position: Step 4 of the workflow
    4. The backend extracts text from the resume

Output feeds directly into ml/keyword_extractor.py.
"""

import os
from docx import Document
from pypdf import PdfReader


class UnsupportedFileTypeError(Exception):
    pass


def extract_text_from_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text).strip()


def extract_text_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables, since some resume templates put
    # skills/experience in table cells rather than plain paragraphs.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs).strip()


def extract_text(file_path: str) -> str:
    """
    Dispatch to the correct extractor based on file extension.
    Raises UnsupportedFileTypeError for anything other than .pdf / .docx.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {ext}. Only .pdf and .docx are supported.")


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        print("Usage: python file_extractor.py <path-to-resume.pdf|.docx>")
    else:
        print(extract_text(sys.argv[1]))
